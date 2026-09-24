# -*- coding: utf-8 -*-
"""[사용자 요청] 라우팅 표를 되살린다.

내가 "그림 5와 겹친다"고 판단해 뺐는데 사용자가 "이거 뿐만 아니라 라우팅표도
필요할듯"이라고 다시 요청했다. 판단이 틀렸다 — 표는 "α가 바뀌면 목적지가 어떻게
달라지나"를, 그림은 "1년 전체가 실제로 어디로 갔나"를 말한다. 다른 질문이다.

수치는 그림 4와 같은 슬롯(h=3499, 한국 출발)에서 본문 3.3절의 점수식으로 다시 뽑았다.
"""
import copy
import re
import shutil

import docx
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt

DOC = "/Users/jongha/Desktop/GitHub/carbon-aware-scheduler/paper/CAST_압축본.docx"
ROWS = [("α", "배정 리전", "탄소", "지연"),
        ("0.0 – 0.4", "한국 (홈)", "351", "0"),
        ("0.5 – 0.9", "캘리포니아", "14", "134"),
        ("1.0", "프랑스", "9", "238")]
W = [1080, 1480, 880, 860]
CAP = ("표 1. 가중치 α 에 따른 배정 변화 — 그림 4와 같은 슬롯이다. 탄소는 gCO₂/kWh, 지연은 "
       "ms. 무릎점 0.508은 가운데 구간에 들어 탄소를 홈 대비 96% 줄이면서 지연은 상한의 "
       "55%만 쓴다.")


def main():
    shutil.copy2(DOC, DOC.replace("CAST_압축본.docx",
                 "versions/2026-09-24/CAST_압축본_라우팅표복원전.docx"))
    d = docx.Document(DOC)
    paras = d.paragraphs
    ai = next(i for i, p in enumerate(paras)
              if p.text.strip().startswith("가중치가 배정을 어떻게 바꾸는지는"))
    anchor = paras[ai]
    src = d.tables[0]._tbl
    cap_proto = next(p for p in paras if re.match(r"^표\s?\d+\.", p.text.strip()))

    t = d.add_table(rows=len(ROWS), cols=4)
    old = t._tbl.find(qn("w:tblPr"))
    if old is not None:
        t._tbl.remove(old)
    t._tbl.insert(0, copy.deepcopy(src.find(qn("w:tblPr"))))
    hdr = src.findall(qn("w:tr"))[0].findall(qn("w:tc"))[0].find(qn("w:tcPr"))
    lay = OxmlElement("w:tblLayout"); lay.set(qn("w:type"), "fixed")
    t._tbl.find(qn("w:tblPr")).append(lay)
    for i, row in enumerate(ROWS):
        for j, v in enumerate(row):
            c = t.cell(i, j)
            if i == 0 and hdr is not None:
                cp = c._tc.find(qn("w:tcPr"))
                if cp is not None:
                    c._tc.remove(cp)
                c._tc.insert(0, copy.deepcopy(hdr))
            tcPr = c._tc.find(qn("w:tcPr"))
            if tcPr is None:
                tcPr = OxmlElement("w:tcPr"); c._tc.insert(0, tcPr)
            e = OxmlElement("w:tcW"); e.set(qn("w:w"), str(W[j])); e.set(qn("w:type"), "dxa")
            tcPr.append(e)
            c.text = ""
            p = c.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            r = p.add_run(str(v)); r.font.size = Pt(7.5); r.bold = (i == 0)
            if j >= 2:
                p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    t._tbl.getparent().remove(t._tbl)

    cp = docx.text.paragraph.Paragraph(copy.deepcopy(cap_proto._p), cap_proto._parent)
    for r in list(cp._p.findall(qn("w:r"))):
        cp._p.remove(r)
    cp.add_run(CAP).font.size = Pt(9)

    body = anchor._p.getparent()
    i = list(body).index(anchor._p) + 1
    body.insert(i, t._tbl)
    body.insert(i + 1, cp._p)
    d.save(DOC)
    print("라우팅 표 복원")

    # 표 번호 재매김
    import renumber_tables
    renumber_tables.main()


if __name__ == "__main__":
    main()
