# -*- coding: utf-8 -*-
"""용량 스윕 서술을 6절(결론)에서 4.3절(용량 제약과 시간 이동)로 옮긴다.

경위: 앞서 f6의 스윕 수치를 넣을 때 "본 연구의 주된 발견은…" 문단 앞에 붙였는데,
그 문단이 결론에 있었다. 그래서 상한 6/12/무제약 민감도가 결론에 실려 있다.
민감도 분석은 결과 절의 소관이다 — 결론은 발견을 말하는 자리다.

옮기면서 둘로 나눈다.
  · 스윕 수치 → 4.3절 끝 (용량 논의의 마무리)
  · 분포 발견(71.9%·91.38%) → 6절에 그대로 남긴다
"""
import re
import shutil
import sys

sys.path.insert(0, "/Users/jongha/Desktop/GitHub/carbon-aware-scheduler/paper/tools")
import docx_edit as D
import docx
from docx.oxml.ns import qn

DOC = "/Users/jongha/Desktop/GitHub/carbon-aware-scheduler/paper/CAST_압축본.docx"

SWEEP = ("유효 상한을 바꾸어 보면 이 결과가 상한값에 얼마나 매여 있는지 드러난다. 상한을 "
         "6건으로 조이면 절감률이 58.05%로 내려가고, 마감이 임박해 미룰 수 없는 작업 때문에 "
         "그 상한 자체를 29,962번 넘긴다. 반대로 상한을 풀면 절감률이 65.80%까지 오르지만, "
         "물리적 상한 12건을 기준으로 다시 세면 위반이 25,992건이다. 본 연구가 쓰는 12건은 "
         "그 사이에서 63.03%·227건에 해당한다.")
CONCL = ("본 연구의 주된 발견은 총계가 아니라 그 분포에 있다. 시간 이동이 만든 절감의 71.9%가 "
         "캘리포니아 한 리전에서 발생하며, 용량 초과가 가장 심한 리전도 캘리포니아다. 그 "
         "리전에서 상한을 넘긴 슬롯의 91.38%는 19시부터 21시까지 세 시간에 몰린다. 반면 "
         "탄소집약도가 종일 평탄한 프랑스는 배정 물량이 가장 많음에도 절감 기여가 2.3%에 "
         "그친다. 시간 이동이 값어치를 만드는 리전과 용량이 그 값어치를 가로막는 리전이 "
         "같으며, 선행 연구가 시간 이동의 예외로 지목한 리전이 바로 그곳이었다.")


def main():
    shutil.copy2(DOC, DOC.replace("CAST_압축본.docx",
                 "versions/2026-09-24/CAST_압축본_스윕이동전.docx"))
    # 1) 결론 문단을 분포 발견만 남기도록 축소
    z, xml = D.load(DOC)
    n0 = len(D.paragraphs(xml))
    xml, o, n = D.replace_text(xml, "유효 상한을 바꾸어 보면", CONCL)
    print(f"  결론 문단 {o} → {n}자 (스윕 서술 분리)")
    assert len(D.paragraphs(xml)) == n0
    D.save(DOC, z, xml, 0)

    # 2) 스윕 서술을 4.3절 끝에 새 문단으로 삽입
    d = docx.Document(DOC)
    paras = d.paragraphs
    # 4.3절 마지막 본문 = 그림 9 캡션 직전 문단
    gi = next(i for i, p in enumerate(paras) if p.text.strip().startswith("그림 9."))
    anchor = paras[gi - 2]          # 그림 문단 앞의 본문
    while not anchor.text.strip():
        gi -= 1
        anchor = paras[gi - 2]
    import copy
    el = copy.deepcopy(anchor._p)
    p = docx.text.paragraph.Paragraph(el, anchor._parent)
    for r in list(p._p.findall(qn("w:r"))):
        p._p.remove(r)
    from docx.shared import Pt
    p.add_run(SWEEP).font.size = Pt(9)
    body = anchor._p.getparent()
    body.insert(list(body).index(anchor._p) + 1, p._p)
    d.save(DOC)
    print(f"  4.3절에 스윕 문단 삽입 (기준: '{anchor.text.strip()[:34]}…' 뒤)")


if __name__ == "__main__":
    main()
