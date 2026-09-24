# -*- coding: utf-8 -*-
"""표 2(선행 정책 비교)의 열 폭과 라벨을 손본다.

기존 표(6행)의 tblPr 을 물려받았더니 첫 열 폭이 그 표 기준이라, 정책명이 긴
새 표에서는 한 칸이 세 줄로 접혀 표가 세 배로 길어졌다. 열 폭을 명시하고
라벨을 짧게 줄인다. 숫자 열은 자릿수가 정해져 있어 좁아도 안 접힌다.
"""
import shutil

import docx
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.shared import Pt

DOC = "/Users/jongha/Desktop/GitHub/carbon-aware-scheduler/paper/CAST_압축본.docx"

LABELS = [
    "정책", "기준(홈 리전 즉시)", "CICS 가상 용량 곡선", "본 연구 — 시간만",
    "본 연구 — 공간만", "본 연구 — 전체", "본 연구 — 시간 무제약",
    "Sukprasert 시간", "CASPER CAP", "Sukprasert 공간+시간",
]
# 단내 폭 4300 dxa 를 나눈다 — 정책명 절반, 숫자 세 열이 나머지
WIDTHS = [1980, 830, 700, 790]


def set_w(cell, w):
    tcPr = cell._tc.find(qn("w:tcPr"))
    if tcPr is None:
        tcPr = OxmlElement("w:tcPr")
        cell._tc.insert(0, tcPr)
    old = tcPr.find(qn("w:tcW"))
    if old is not None:
        tcPr.remove(old)
    e = OxmlElement("w:tcW")
    e.set(qn("w:w"), str(w))
    e.set(qn("w:type"), "dxa")
    tcPr.append(e)


def main():
    shutil.copy2(DOC, DOC.replace("CAST_압축본.docx",
                 "versions/2026-09-24/CAST_압축본_표폭수정전.docx"))
    d = docx.Document(DOC)
    t = [x for x in d.tables if len(x.rows) == 10][0]

    # 고정 레이아웃 — 내용에 따라 열이 흔들리지 않게
    tblPr = t._tbl.find(qn("w:tblPr"))
    lay = OxmlElement("w:tblLayout")
    lay.set(qn("w:type"), "fixed")
    tblPr.append(lay)

    for i, row in enumerate(t.rows):
        for j, cell in enumerate(row.cells):
            set_w(cell, WIDTHS[j])
            if j == 0:
                p = cell.paragraphs[0]
                for r in p.runs:
                    r.text = ""
                r = p.runs[0] if p.runs else p.add_run()
                r.text = LABELS[i]
                r.font.size = Pt(7.5)
                r.bold = (i == 0)

    # 캡션의 K_r=12 표기를 읽기 쉽게
    for p in d.paragraphs:
        if p.text.startswith("표 2. 선행 정책"):
            for r in p.runs:
                r.text = r.text.replace("유효 상한 K_r=12을", "유효 상한 12건을")
    d.save(DOC)
    print("열 폭 고정 + 라벨 축약 완료")


if __name__ == "__main__":
    main()
