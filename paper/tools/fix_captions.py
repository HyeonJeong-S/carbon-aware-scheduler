# -*- coding: utf-8 -*-
"""캡션 정리 — 폰트 통일과 표 번호 중복 해결 (사용자 지적 4·12번).

사용자: "각각 그림 설명의 폰트가 8, 10 등으로 다 다름"
실측 결과 9.0 / 8.0 / 7.5 세 가지가 섞여 있었다. 이 문서의 관례는 9.0 이다.

그리고 점검 중에 더 큰 것이 나왔다 — **표 번호가 중복이다.** 선행 정책 비교표를
"표 2"로 붙였는데 이미 "표 2. 지연 등급별 시간 이동 절감 기여"가 있었다.
등장 순서대로 다시 매긴다: 1 단계별 → 2 지연 등급별 → 3 사후 강제 → 4 선행 정책.
"""
import re
import shutil

import docx
from docx.shared import Pt

DOC = "/Users/jongha/Desktop/GitHub/carbon-aware-scheduler/paper/CAST_압축본.docx"
CAP_PT = 9.0


def main():
    shutil.copy2(DOC, DOC.replace("CAST_압축본.docx",
                 "versions/2026-09-24/CAST_압축본_캡션정리전.docx"))
    d = docx.Document(DOC)

    # 1) 선행 정책 비교표 → 표 4 (본문 참조도 함께)
    for p in d.paragraphs:
        t = p.text.strip()
        if t.startswith("표 2. 선행 정책을 원문"):
            for r in p.runs:
                if "표 2." in r.text:
                    r.text = r.text.replace("표 2.", "표 4.", 1)
                    break
            print("  표 2(선행 정책) → 표 4")
        elif "표 N" in t or "표 2와 같다" in t:
            print(f"  [확인 필요] {t[:50]}")

    # 본문에서 선행 정책 표를 가리키는 곳
    for p in d.paragraphs:
        if "결과는 표 2와 같다" in p.text:
            for r in p.runs:
                if "표 2" in r.text:
                    r.text = r.text.replace("표 2", "표 4")
            print("  본문 참조 → 표 4")

    # 2) 캡션 폰트 통일
    n = 0
    for p in d.paragraphs:
        if re.match(r"^(그림|표)\s?\d+\.", p.text.strip()):
            for r in p.runs:
                if r.font.size is None or abs(r.font.size.pt - CAP_PT) > 0.01:
                    r.font.size = Pt(CAP_PT)
                    n += 1
    print(f"  캡션 run {n}개 → {CAP_PT}pt")

    d.save(DOC)

    # 검증
    d2 = docx.Document(DOC)
    print("\n=== 캡션 최종 ===")
    for p in d2.paragraphs:
        t = p.text.strip()
        if re.match(r"^(그림|표)\s?\d+\.", t):
            s = sorted({(r.font.size.pt if r.font.size else None)
                        for r in p.runs if r.text.strip()})
            print(f"  {str(s):<10} {t[:56]}")


if __name__ == "__main__":
    main()
