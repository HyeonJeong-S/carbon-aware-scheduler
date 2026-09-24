# -*- coding: utf-8 -*-
"""[16] 공간 이동 라우팅 그림을 3.3절에 넣는다.

사용자: "공간이동은 다른 논문들이 공간이동을 하는 그런 아키텍처를 따와서 그리는 게
낫지 않아??"  참고 문헌(CASPER Fig.3a)이 리전 간 화살표로 재배치를 보이는 관례를
따르되, 도식이 아니라 실측 배정 행렬(146,000건)을 그렸다.

표 1(α → 배정 표)은 그대로 둔다. 표는 "α를 바꾸면 목적지가 어떻게 달라지는가"를,
그림은 "그래서 1년치 전체가 실제로 어디로 갔는가"를 말한다 — 다른 질문에 답한다.

넣는 자리: 표 1 캡션 뒤(무릎점 설명 직후). 그림 번호는 표 1 다음이므로 그림 5가
되고, 이후 번호를 한 칸씩 민다.
"""
import copy
import re
import shutil

import docx
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Pt

DOC = "/Users/jongha/Desktop/GitHub/carbon-aware-scheduler/paper/CAST_압축본.docx"
FIG = "/Users/jongha/Desktop/GitHub/carbon-aware-scheduler/paper/diagram/fig4_routing.png"
CAP = ("그림 5. 공간 이동이 실제로 만든 리전 간 흐름 — 무릎점 자동 선택으로 1년치 146,000건을 "
       "배정한 결과다. 위가 출발, 아래가 도착이며 선 굵기가 옮긴 작업 수, 회색 세로선이 홈에 "
       "남은 몫이다. 프랑스와 캘리포니아가 62.3%를 받고 인도는 1.9%만 남는다. 이 쏠림이 "
       "4.3절에서 다루는 용량 문제의 원인이다.")


def set_text(p, new):
    runs = p._p.findall(qn("w:r"))
    for r in runs[1:]:
        p._p.remove(r)
    for e in runs[0].findall(qn("w:t")):
        runs[0].remove(e)
    el = runs[0].makeelement(qn("w:t"), {})
    el.set(qn("xml:space"), "preserve")
    el.text = new
    runs[0].append(el)


def main():
    shutil.copy2(DOC, DOC.replace("CAST_압축본.docx",
                 "versions/2026-09-24/CAST_압축본_라우팅그림전.docx"))
    d = docx.Document(DOC)

    # 그림 5 이상을 한 칸씩 (큰 번호부터)
    caps = {}
    for p in d.paragraphs:
        m = re.match(r"^그림\s?(\d+)\.", p.text.strip())
        if m:
            caps[int(m.group(1))] = p
    for n in sorted((k for k in caps if k >= 5), reverse=True):
        set_text(caps[n], re.sub(r"^그림\s?\d+\.", f"그림 {n+1}.", caps[n].text.strip()))
        print(f"  그림 {n} → 그림 {n+1}")
    for n in sorted((k for k in caps if k >= 5), reverse=True):
        for p in d.paragraphs:
            if re.match(r"^그림\s?\d+\.", p.text.strip()):
                continue
            if f"그림 {n}" in p.text:
                for r in p.runs:
                    if f"그림 {n}" in r.text:
                        r.text = r.text.replace(f"그림 {n}", f"그림 {n+1}")
                print(f"  본문 참조 그림 {n} → {n+1}")

    # 표 1 캡션 뒤에 삽입
    paras = d.paragraphs
    ai = next(i for i, p in enumerate(paras)
              if p.text.strip().startswith("표 1. 가중치 α"))
    anchor = paras[ai]
    proto = next(p for p in paras if p.text.strip().startswith("그림 4."))

    fp = d.add_paragraph()
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fp.add_run().add_picture(FIG, width=Pt(205))
    cp = docx.text.paragraph.Paragraph(copy.deepcopy(proto._p), proto._parent)
    for r in list(cp._p.findall(qn("w:r"))):
        cp._p.remove(r)
    cp.add_run(CAP).font.size = Pt(9)

    body = anchor._p.getparent()
    i = list(body).index(anchor._p)
    fp._p.getparent().remove(fp._p)
    body.insert(i + 1, fp._p)
    body.insert(i + 2, cp._p)

    d.save(DOC)
    d2 = docx.Document(DOC)
    print(f"\n그림 {sum(1 for p in d2.paragraphs if 'graphicData' in p._p.xml)}장")
    for p in d2.paragraphs:
        if re.match(r"^(그림|표)\s?\d+\.", p.text.strip()):
            print("  ", p.text.strip()[:54])


if __name__ == "__main__":
    main()
