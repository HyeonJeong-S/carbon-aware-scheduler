# -*- coding: utf-8 -*-
"""본문 기호를 수식과 같은 글씨체로 — Cambria Math 이탤릭 + 진짜 첨자.

왜 필요한가 (2026-09-24 사용자 메모): "ar이랑 Pslack을 할 때 그 우리가 공식에
적은 폰트로 넣고싶어."  본편 식은 OMML 이라 Cambria Math 로 조판되는데 본문의
같은 기호는 본문 서체 평문이라 한 문장 안에서 글씨가 갈렸다.

**사고에서 배운 것**: run 을 통째로 갈아끼우는 작업은 그 문단에 달린 Word 메모의
`<w:commentReference>` 를 품은 run 까지 지운다. 범위 표식(commentRangeStart/End)은
w:r 이 아니라 살아남지만 참조가 사라져 **어디에도 안 달린 떠돌이 메모**가 된다.
실제로 메모 둘을 그렇게 잃었다(내용은 문서/메모_받은것_*.txt 에 보존돼 있었다).
그래서 여기서는 메모 표식을 먼저 뽑아 두었다가 새 run 앞뒤로 도로 끼운다.

실행: ./.venv/bin/python paper/tools/mathify.py paper/CAST_압축본.docx [--dry]
"""
import copy
import re
import shutil
import sys
from datetime import datetime

import docx
from docx.oxml.ns import qn

MATH = "Cambria Math"

# 첨자가 붙는 기호 — (밑, 첨자) 짝. 긴 것부터 와야 P_slack 이 P_s 로 안 잘린다.
SYM = [("P", "slack"), ("cap", "r"), ("t", "early"), ("t", "late"), ("x", "jr"),
       ("S", "r"), ("F", "r"), ("P", "r"), ("R", "r"), ("Q", "t"), ("W", "j"),
       ("a", "r"), ("z", "j"), ("K", "r"), ("n", "r"), ("u", "j"), ("o", "j"),
       ("s", "j"), ("d", "j"), ("D", "j"), ("L", "j"), ("k", "j"), ("τ", "j"),
       ("C", "r")]
# 두 가지 꼴을 다 받는다 — 밑줄이 살아 있는 "a_r"(원고 원문)과 워드가 저장하며
# 첨자를 평평하게 눌러 버린 "ar"(사용자가 한 번 열었다 저장한 파일). 둘 다 실제로
# 겪었다.
SUB_RE = re.compile(
    "(?<![A-Za-z가-힣])("
    + "|".join(re.escape(b + "_" + s) for b, s in SYM) + "|"
    + "|".join(re.escape(b + s) for b, s in SYM)
    + ")(?![A-Za-z])")
LOOK = {}
for _b, _s in SYM:
    LOOK[_b + _s] = (_b, _s)
    LOOK[_b + "_" + _s] = (_b, _s)

# 맨 글자 변수. 앞뒤가 식별자·%·하이픈이면 아니다(E-mail, 0.11%p, 저자 이니셜).
BARE_RE = re.compile(r"(?<![A-Za-z0-9α-ωΑ-Ω._/@%-])([a-zA-Zα-ωΑ-Ω])"
                     r"(?![A-Za-z0-9α-ωΑ-Ω._/@-])")

SKIP_PREFIX = ("Algorithm 1",)


def _korean(s):
    return sum(1 for c in s if "가" <= c <= "힣") / max(1, len(s))


def _targets(p):
    """이 문단에서 바꿀 (시작, 끝, 밑, 첨자) 목록.

    이미 Cambria Math 로 바뀐 자리는 건너뛴다 — 두 번 돌려도 같은 결과여야
    한다(멱등). 그러지 않으면 돌릴 때마다 run 이 쪼개졌다 붙었다 하면서
    메모·첨자 같은 남은 서식을 잃을 기회만 늘어난다.
    """
    t = p.text
    done = set()
    pos = 0
    for r in p.runs:
        if r.font.name == MATH:
            done |= set(range(pos, pos + len(r.text)))
        pos += len(r.text)

    spans = []
    for m in SUB_RE.finditer(t):
        if m.start() in done:
            continue
        b, s = LOOK[m.group(1)]
        spans.append((m.start(), m.end(), b, s))
    taken = done | {i for a, b_, _, _ in spans for i in range(a, b_)}

    if _korean(t) >= 0.25 and "@" not in t and "E-mail" not in t:
        for m in BARE_RE.finditer(t):
            if m.start() not in taken:
                spans.append((m.start(), m.end(), m.group(1), ""))
    return sorted(spans)


def mathify_par(p):
    if "<m:oMath" in p._p.xml or p.text.startswith(SKIP_PREFIX):
        return 0
    if p.style.name.startswith("Heading") or not p.runs:
        return 0
    spans = _targets(p)
    if not spans:
        return 0

    t = p.text
    # 글자 → 원래 run 지도. 바꾸지 않는 구간의 서식(첨자 등)을 그대로 되살린다.
    owner = [r for r in p.runs for _ in r.text]
    base_rpr = p.runs[0]._r.find(qn("w:rPr"))

    frag = p._p.xml
    starts = "".join(re.findall(r"<w:commentRangeStart\b[^>]*/>", frag))
    ends = "".join(re.findall(r"<w:commentRangeEnd\b[^>]*/>", frag))
    refs = "".join(re.findall(
        r"<w:r\b(?:(?!</w:r>).)*?<w:commentReference\b[^>]*/>(?:(?!</w:r>).)*?</w:r>",
        frag, re.DOTALL))

    for r in list(p._p.findall(qn("w:r"))):
        p._p.remove(r)
    for e in p._p.findall(qn("w:commentRangeStart")) + p._p.findall(qn("w:commentRangeEnd")):
        p._p.remove(e)

    def add(text, kind, src=None):
        if not text:
            return
        r = p._p.makeelement(qn("w:r"), {})
        src_rpr = src if src is not None else base_rpr
        c = copy.deepcopy(src_rpr) if src_rpr is not None else r.makeelement(qn("w:rPr"), {})
        if kind in ("m", "ms"):
            for e in c.findall(qn("w:rFonts")):
                c.remove(e)
            c.insert(0, c.makeelement(qn("w:rFonts"),
                                      {qn("w:ascii"): MATH, qn("w:hAnsi"): MATH}))
            c.append(c.makeelement(qn("w:i"), {}))
            if kind == "ms":
                c.append(c.makeelement(qn("w:vertAlign"), {qn("w:val"): "subscript"}))
        r.append(c)
        w = r.makeelement(qn("w:t"), {qn("xml:space"): "preserve"})
        w.text = text
        r.append(w)
        p._p.append(r)

    for frag_xml in re.findall(r"<w:commentRangeStart\b[^>]*/>", starts):
        p._p.append(_el(frag_xml))

    last = 0
    for a, b_, base, sub in spans:
        # 안 바뀌는 구간은 원래 서식대로 쪼개 쓴다
        i = last
        while i < a:
            j = i
            while j < a and owner[j] is owner[i]:
                j += 1
            add(t[i:j], "n", owner[i]._r.find(qn("w:rPr")))
            i = j
        add(base, "m")
        add(sub, "ms")
        last = b_
    i = last
    while i < len(t):
        j = i
        while j < len(t) and owner[j] is owner[i]:
            j += 1
        add(t[i:j], "n", owner[i]._r.find(qn("w:rPr")))
        i = j

    for frag_xml in re.findall(r"<w:commentRangeEnd\b[^>]*/>", ends):
        p._p.append(_el(frag_xml))
    for frag_xml in re.findall(
            r"<w:r\b(?:(?!</w:r>).)*?<w:commentReference\b[^>]*/>(?:(?!</w:r>).)*?</w:r>",
            refs, re.DOTALL):
        p._p.append(_el(frag_xml))
    return len(spans)


def _el(xml_frag):
    NS = ('xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" '
          'xmlns:w14="http://schemas.microsoft.com/office/word/2010/wordml"')
    return docx.oxml.parse_xml(f"<w:body {NS}>{xml_frag}</w:body>")[0]


def main():
    path = sys.argv[1]
    dry = "--dry" in sys.argv
    if not dry:
        snap = path[:-5] + f"_{datetime.now():%Y%m%d_%H%M%S}_before_mathify.docx"
        snap = snap.replace("/paper/", "/paper/versions/")
        shutil.copy2(path, snap)
        print("스냅샷:", snap.split("/")[-1])
    d = docx.Document(path)
    before = [p.text for p in d.paragraphs]
    n = sum(mathify_par(p) for p in d.paragraphs)
    after = [p.text for p in d.paragraphs]
    # "a_r" 을 첨자로 올리면 밑줄이 사라지므로 글자가 그대로일 수 없다(2026-09-24
    # 이 검사가 잘못 걸려 저장을 막았다). 밑줄을 뺀 뒤 비교하면 진짜 글자 손실만
    # 잡힌다 — 파일 이름의 밑줄은 양쪽에서 똑같이 빠지므로 영향이 없다.
    assert [x.replace("_", "") for x in before] == [x.replace("_", "") for x in after], \
        "본문 글자가 바뀌었다 — 저장하지 않는다"
    if dry:
        print(f"[시험] {n}곳을 바꿀 것이다")
        return
    d.save(path)
    chk = docx.Document(path)
    import zipfile
    x = zipfile.ZipFile(path).read("word/document.xml").decode()
    print(f"{n}곳 적용 · 문단 {len(chk.paragraphs)} · 그림 {x.count('<w:drawing')} "
          f"· 메모 참조 {x.count('commentReference')}")


if __name__ == "__main__":
    main()
