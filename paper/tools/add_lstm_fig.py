# -*- coding: utf-8 -*-
"""LSTM 예측 그림을 그림 2로 넣고 뒤 번호를 한 칸씩 민다 (2026-09-24).

사용자 지시: "전체, lstm, 로드밸런서, 스케줄러 이런 그림 4개는 그려야 사람들이
이해할 수 있을거야." 압축본에서 §5.3(예측)을 통째로 잘라내면서 예측 그림도 같이
빠져 있었다. 그림이 잘려나간 절의 내용을 대신하므로 글을 늘리지 않고 복원된다.

같이 처리하는 것
 · 그림 2·4 캡션 교체 — 그림이 개념도에서 실측 자료 그림으로 바뀌어 옛 캡션이
   더 이상 그림을 설명하지 않는다.
 · "그림 N은 …이다" 꼴 문단 3개 삭제 — 캡션이 이미 말한 것을 되풀이하고,
   그중 둘은 바뀐 그림을 **틀리게** 설명한다(입력 상자·처리 순서 운운).
"""
import copy
import shutil

import docx
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Pt

REPO = "/Users/jongha/Desktop/GitHub/carbon-aware-scheduler/"
DOC = REPO + "paper/CAST_압축본.docx"
FIG = REPO + "paper/diagram/fig2_forecast.png"

NEW_CAP = {
    # 새 그림 번호 기준
    3: ("그림 3. 공간 이동이 맞바꾸는 두 축 — 리전 간 격차가 큰 한 슬롯에서 한국이 출발지일 "
        "때 여덟 리전의 위치다. 홈에 두면 351 gCO₂/kWh이고, 134 ms를 쓰면 337이 줄지만, "
        "거기서 104 ms를 더 써도 5밖에 줄지 않는다. 무릎점은 이 꺾이는 자리를 슬롯마다 "
        "다시 고른다."),
    5: ("그림 5. 용량이 시간 이동을 막는 방식 — 캘리포니아의 한 날 실측이다. 위는 시각별 "
        "탄소집약도, 아래는 같은 시각의 동시 실행 수다. 탄소가 낮아지는 저녁 시간대가 그대로 "
        "유효 상한에 닿아 있어, 옮기고 싶은 자리가 이미 차 있다. 상한을 넘는 막대는 마감이 "
        "임박해 미룰 수 없는 작업이다."),
}
DROP = ["그림 1은 세 단계의 연결과",
        "그림 2는 예측·지연 행렬·직전 배정이라는",
        "그림 4는 그림 2가 정한 리전 배정을"]
LSTM_CAP = ("그림 2. 리전별 탄소집약도 예측 — 지난 168시간을 받아 향후 24시간을 낸다. "
            "캘리포니아의 한 구간이며 평균 절대 오차는 16.4 gCO₂/kWh다. 리전마다 독립된 "
            "2층 LSTM(은닉 64)을 쓰고, 배치 결정에는 이 예측값을, 배출량 산정에는 실측값을 쓴다.")


def main():
    shutil.copy2(DOC, REPO + "paper/versions/2026-09-24/CAST_압축본_LSTM그림전.docx")
    d = docx.Document(DOC)
    paras = d.paragraphs

    # ── 1) 바뀐 그림을 틀리게 설명하는 문단 삭제 ──
    dropped = 0
    for head in DROP:
        for p in list(d.paragraphs):
            if p.text.strip().startswith(head):
                p._p.getparent().remove(p._p)
                dropped += 1
                print(f"  삭제  {head[:28]}…")
                break

    # ── 2) 캡션 번호 밀기 (뒤에서부터 — 7→8, 6→7, …, 2→3) ──
    caps = {}
    for p in d.paragraphs:
        t = p.text.strip()
        if t.startswith("그림 ") and ". " in t[:8]:
            n = int(t.split(".")[0].replace("그림", "").strip())
            caps[n] = p
    for n in sorted((k for k in caps if k >= 2), reverse=True):
        p = caps[n]
        new_n = n + 1
        if new_n in NEW_CAP:
            for r in p.runs:
                r.text = ""
            (p.runs[0] if p.runs else p.add_run()).text = NEW_CAP[new_n]
            print(f"  캡션  그림 {n} → 그림 {new_n} (내용 교체)")
        else:
            done = False
            for r in p.runs:
                if not done and f"그림 {n}" in r.text:
                    r.text = r.text.replace(f"그림 {n}", f"그림 {new_n}", 1)
                    done = True
            if not done:                      # run 이 쪼개진 경우
                txt = p.text.replace(f"그림 {n}", f"그림 {new_n}", 1)
                for r in p.runs:
                    r.text = ""
                p.runs[0].text = txt
            print(f"  캡션  그림 {n} → 그림 {new_n}")

    # ── 3) 본문 안의 그림 참조도 함께 밀기 ──
    for p in d.paragraphs:
        t = p.text
        if "그림" not in t or t.strip().startswith("그림 "):
            continue
        for n in sorted((k for k in caps if k >= 2), reverse=True):
            if f"그림 {n}" in t:
                for r in p.runs:
                    if f"그림 {n}" in r.text:
                        r.text = r.text.replace(f"그림 {n}", f"그림 {n+1}")
                print(f"  본문  그림 {n} → 그림 {n+1}  | {p.text[:46]}")

    # ── 4) 새 그림 2를 그림 1 캡션 바로 뒤에 넣는다 ──
    anchor = None
    for p in d.paragraphs:
        if p.text.strip().startswith("그림 1. CAST의 전체 파이프라인"):
            anchor = p
            break
    assert anchor is not None, "그림 1 캡션을 찾지 못했다"
    proto_cap = anchor

    fp = d.add_paragraph()
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fp.add_run().add_picture(FIG, width=Pt(205))
    cp = docx.text.paragraph.Paragraph(copy.deepcopy(proto_cap._p), proto_cap._parent)
    for r in list(cp._p.findall(qn("w:r"))):
        cp._p.remove(r)
    r = cp.add_run(LSTM_CAP)
    r.font.size = Pt(8)

    body = anchor._p.getparent()
    i = list(body).index(anchor._p)
    fp._p.getparent().remove(fp._p)
    body.insert(i + 1, fp._p)
    body.insert(i + 2, cp._p)

    d.save(DOC)
    d2 = docx.Document(DOC)
    n_fig = sum(1 for p in d2.paragraphs if "graphicData" in p._p.xml)
    print(f"\n그림 {n_fig}개, 문단 {len(d2.paragraphs)}개 (설명 문단 {dropped}개 삭제)")


if __name__ == "__main__":
    main()
