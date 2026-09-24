# -*- coding: utf-8 -*-
"""사용자 피드백 일괄 반영 1차 — 6·12·13번 + 캡션 마침표 (2026-09-24 야간).

[6]  3.4절 "이 단계의 가중치를 무릎점으로 결정하지 않는 이유는…" 문단 삭제
     사용자: "이 이야기는 안 해도 될 거 같지 않아?" — 동의한다. 왜 안 했는지를
     설명하는 문단은 독자에게 새 정보를 주지 않고 의심만 남긴다.
[12] 표 4 의 절감률 열에서 "%"가 줄바꿈된다 — 열 폭 재배분
[13] 그림 8 캡션이 무슨 그림인지 먼저 말하지 않아 "선행 정책 비교 그래프가
     있으면 좋겠다"는 요청을 또 받았다. 캡션 첫 구절을 바꾼다.
(+)  그림 6 캡션 끝 마침표 누락 (인용포함본에는 있음)
"""
import re
import shutil
import sys

import docx
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt

sys.path.insert(0, "/Users/jongha/Desktop/GitHub/carbon-aware-scheduler/paper/tools")
import docx_edit as D

DOC = "/Users/jongha/Desktop/GitHub/carbon-aware-scheduler/paper/CAST_압축본.docx"


def text_edits():
    z, xml = D.load(DOC)
    n0 = len(D.paragraphs(xml))

    # [6] 무릎점 미사용 이유 문단 삭제
    xml, o = D.delete(xml, "이 단계의 가중치를 무릎점으로 결정하지 않는 이유는")
    print(f"  [6]  삭제 {o}자 — 무릎점 미사용 이유 문단")

    # [13] 그림 8 캡션 — 무슨 그림인지 먼저
    xml, a, b = D.replace_text(xml, "그림 8. 탄소 절감률과 용량 상한 위반",
        "그림 8. 선행 정책 재현 결과 비교 (표 4) — 세로축은 탄소 절감률, 가로축은 같은 "
        "기준으로 센 용량 상한 위반 건수다. ★ 본 연구, ○ 선행 정책. 왼쪽 위가 좋은 "
        "자리이며, 본 연구를 앞서는 정책은 모두 오른쪽 바깥, 곧 더 큰 과부하를 대가로 "
        "치른 자리에 있다.")
    print(f"  [13] 캡션 {a} → {b}자 — 무슨 그림인지 먼저 밝힘")

    # (+) 그림 6 캡션 마침표
    xml, a, b = D.replace_text(xml, "그림 6. 다섯 방식의 총 배출량 비교",
        "그림 6. 다섯 방식의 총 배출량 비교(표 1) — (a)전체, (b)①을 제외하고 x축을 확대한 "
        "것이며 점선은 ④ 값이다. 온라인 용량 인지(④, 10,805.0 kg, −63.03%)는 아무것도 안 "
        "했을 때(①)와, 용량 제약을 전혀 고려하지 않아 실현 불가능한 반사실 상한(③, "
        "9,958.2 kg) 사이에 위치한다.")
    print(f"  (+)  그림 6 캡션 마침표 복원 ({a} → {b}자)")

    assert len(D.paragraphs(xml)) == n0 - 1
    assert xml.count("<w:drawing>") == 8
    D.save(DOC, z, xml, -1)


def table_widths():
    """[12] 표 4 절감률 열이 좁아 '%'가 줄바꿈된다."""
    d = docx.Document(DOC)
    t = [x for x in d.tables if len(x.rows) == 10][0]
    W = [1850, 880, 810, 760]          # 정책 / 총배출 / 절감률 / 위반
    for row in t.rows:
        for j, cell in enumerate(row.cells):
            tcPr = cell._tc.find(qn("w:tcPr"))
            if tcPr is None:
                tcPr = OxmlElement("w:tcPr")
                cell._tc.insert(0, tcPr)
            old = tcPr.find(qn("w:tcW"))
            if old is not None:
                tcPr.remove(old)
            e = OxmlElement("w:tcW")
            e.set(qn("w:w"), str(W[j]))
            e.set(qn("w:type"), "dxa")
            tcPr.append(e)
            for p in cell.paragraphs:
                for r in p.runs:
                    r.font.size = Pt(7.0)      # 7.5 → 7.0
    d.save(DOC)
    print(f"  [12] 표 4 열 폭 {W} · 글씨 7.0pt")


if __name__ == "__main__":
    shutil.copy2(DOC, DOC.replace("CAST_압축본.docx",
                 "versions/2026-09-24/CAST_압축본_batch1전.docx"))
    text_edits()
    table_widths()
