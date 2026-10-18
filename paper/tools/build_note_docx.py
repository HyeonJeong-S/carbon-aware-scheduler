# -*- coding: utf-8 -*-
"""짧은 정리 문서(스토리·면담 안건·피드백 대응표) txt → A4 docx. 10-19.
    ./.venv/bin/python paper/tools/build_note_docx.py <입력.txt> <출력.docx>
표기: "# 제목" · "## 절" · "■ 묶음 제목" · 빈 줄로 단락 · "· " 항목 · "| a | b |" 표(첫 행이 머리) ·
      "[그림: 파일이름 | 캡션]"(paper/diagram/<파일>.png) · "=== 쪽 ===" 쪽 나눔 · "#" 로 시작하고 바로 글자가 붙은 줄은 메모(버림).
PDF 는 Word 로(paper/tools/word_to_pdf.applescript — 경로 제약 때문에 paper/CAST_JOK_v2.docx 자리를 빌려 쓴다)."""
import os
import re
import sys

import docx
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_BREAK
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor

ROOT = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
DIAG = os.path.join(ROOT, "paper", "diagram")
FONT = "맑은 고딕"


def _font(run, size=None, bold=None, color=None):
    run.font.name = FONT
    rpr = run._element.get_or_add_rPr()
    rf = rpr.find(qn("w:rFonts"))
    if rf is None:
        rf = rpr.makeelement(qn("w:rFonts"), {}); rpr.append(rf)
    for a in ("ascii", "hAnsi", "eastAsia", "cs"):
        rf.set(qn("w:" + a), FONT)
    if size: run.font.size = Pt(size)
    if bold is not None: run.font.bold = bold
    if color: run.font.color.rgb = RGBColor(*color)


def _para(doc, text, size=10, bold=False, space_after=3, indent=None, color=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(space_after); p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.line_spacing = 1.25
    if indent is not None:
        p.paragraph_format.left_indent = Mm(indent); p.paragraph_format.first_line_indent = Mm(-3)
    for k, part in enumerate(re.split(r"\*\*(.+?)\*\*", text)):   # **굵게**
        if part: _font(p.add_run(part), size, bold or k % 2 == 1, color)
    return p


def _table(doc, rows):
    ncol = max(len(r) for r in rows)
    t = doc.add_table(rows=len(rows), cols=ncol); t.style = "Table Grid"; t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, r in enumerate(rows):
        for j in range(ncol):
            c = t.cell(i, j); c.text = ""
            p = c.paragraphs[0]; p.paragraph_format.space_after = Pt(0); p.paragraph_format.line_spacing = 1.1
            _font(p.add_run(r[j] if j < len(r) else ""), 8.5, i == 0)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def build(src, out):
    doc = docx.Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Mm(210), Mm(297)
    sec.left_margin = sec.right_margin = Mm(20); sec.top_margin = sec.bottom_margin = Mm(18)
    lines = open(src, encoding="utf-8").read().split("\n")
    buf, rows = [], []

    def flush():
        nonlocal buf
        if buf: _para(doc, " ".join(x.strip() for x in buf)); buf = []

    def flush_rows():
        nonlocal rows
        if rows: _table(doc, rows); rows = []

    for ln in lines:
        s = ln.rstrip()
        if re.match(r"^#[^#\s]", s):          # 메모
            continue
        if s.startswith("|"):
            flush(); rows.append([c.strip() for c in s.strip("|").split("|")]); continue
        flush_rows()
        if not s.strip():
            flush(); continue
        if s.startswith("# "):
            flush(); _para(doc, s[2:], 15, True, 8); continue
        if s.startswith("## "):
            flush(); _para(doc, s[3:], 12, True, 4, color=(0x1f, 0x3a, 0x5f)); continue
        if s.startswith("■ "):
            flush(); _para(doc, s[2:], 10.5, True, 2); continue
        if s.startswith("· "):
            flush(); _para(doc, "· " + s[2:], 10, indent=5, space_after=1); continue
        m = re.match(r"^\[그림:\s*([^|\]]+?)\s*\|\s*(.+)\]$", s)
        if m:
            flush()
            path = os.path.join(DIAG, m.group(1) + ".png")
            if os.path.exists(path):
                doc.add_picture(path, width=Mm(110)); doc.paragraphs[-1].alignment = 1
            _para(doc, m.group(2), 8.5, color=(0x44, 0x44, 0x44)); continue
        if s.strip() == "=== 쪽 ===":
            flush(); doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE); continue
        buf.append(s)
    flush(); flush_rows()
    cp = doc.core_properties; cp.author = ""; cp.last_modified_by = ""
    doc.save(out); print("wrote", out)


if __name__ == "__main__":
    build(sys.argv[1], sys.argv[2])
