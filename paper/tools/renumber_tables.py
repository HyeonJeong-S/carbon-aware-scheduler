# -*- coding: utf-8 -*-
"""표 번호를 등장 순서대로 다시 매긴다 — 두 번째 충돌이라 자동화한다.

오늘만 두 번 충돌했다(벤치마크 표를 표 2로, 라우팅 표를 또 표 2로). 표를 새로
끼울 때마다 뒤 번호를 손으로 미는 방식은 반드시 또 틀린다. 문서 순서대로
"표 N." 을 다시 부여하고, 본문의 "표 N" 참조도 같은 사상으로 바꾼다.

참조를 옮길 때 주의: 낡은 번호와 새 번호가 겹칠 수 있으므로 **한 번에** 치환해야
한다(순차 치환하면 방금 만든 번호를 다시 집는다 — 식 번호에서 당한 그 사고).
"""
import re
import shutil

import docx
from docx.oxml.ns import qn

DOC = "/Users/jongha/Desktop/GitHub/carbon-aware-scheduler/paper/CAST_압축본.docx"


def set_text(p, new):
    runs = p._p.findall(qn("w:r"))
    if not runs:
        return
    for r in runs[1:]:
        p._p.remove(r)
    for t in runs[0].findall(qn("w:t")):
        runs[0].remove(t)
    el = runs[0].makeelement(qn("w:t"), {})
    el.set(qn("xml:space"), "preserve")
    el.text = new
    runs[0].append(el)


def main():
    shutil.copy2(DOC, DOC.replace("CAST_압축본.docx",
                 "versions/2026-09-24/CAST_압축본_표번호전.docx"))
    d = docx.Document(DOC)

    caps = [p for p in d.paragraphs if re.match(r"^표\s?\d+\.", p.text.strip())]
    mapping = {}
    for new, p in enumerate(caps, 1):
        old = int(re.match(r"^표\s?(\d+)\.", p.text.strip()).group(1))
        mapping[old] = new
        set_text(p, re.sub(r"^표\s?\d+\.", f"표 {new}.", p.text.strip()))
        print(f"  표 {old} → 표 {new}  {p.text.strip()[:44]}")

    # 본문 참조 — 한 번에 치환(순차로 하면 방금 만든 번호를 다시 집는다)
    pat = re.compile(r"표\s?(\d+)")
    n = 0
    for p in d.paragraphs:
        t = p.text
        if "표" not in t or re.match(r"^표\s?\d+\.", t.strip()):
            continue
        new_t = pat.sub(lambda m: f"표 {mapping.get(int(m.group(1)), m.group(1))}", t)
        if new_t != t:
            set_text(p, new_t)
            n += 1
            print(f"  참조: {t[:44]}… → {new_t[:44]}…")
    print(f"본문 참조 {n}문단 갱신")
    d.save(DOC)

    d2 = docx.Document(DOC)
    print("\n=== 최종 표 순서 ===")
    for p in d2.paragraphs:
        if re.match(r"^표\s?\d+\.", p.text.strip()):
            print("  ", p.text.strip()[:62])


if __name__ == "__main__":
    main()
