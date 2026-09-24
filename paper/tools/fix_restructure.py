# -*- coding: utf-8 -*-
"""재배치 뒤처리 — 앞 스크립트가 놓친 네 가지.

1. 절 제목이 번호만 바뀌고 내용이 그대로였다 ("2. 관련 연구와 그 한계",
   "3. 제안하는 방법"). 사용자가 준 형식은 Related Work / CAST 다.
2. "6.4절"이 세 곳 남았다. 그중 하나는 Sukprasert 논문의 절이라 건드리면 안 된다.
3. Word 가 run 을 어절 단위로 쪼개 "5.6" 과 "절" 이 따로 저장된 곳이 있어
   run 단위 치환이 먹지 않았다 — 문단 평문에서 찾아 run 을 다시 조립한다.
4. 빨간 글씨가 표 안에 남아 있었다 (앞 스크립트는 본문 문단만 훑었다).
"""
import re
import shutil

import docx
from docx.oxml.ns import qn

DOC = "/Users/jongha/Desktop/GitHub/carbon-aware-scheduler/paper/CAST_압축본.docx"

TITLE = {"2. 관련 연구와 그 한계": "2. 관련 연구",
         "3. 제안하는 방법": "3. CAST",
         "4. 검증": "4. 실험과 분석"}
XREF = {"5.6절": "3.5절", "5.5절": "3.4절", "5.4절": "3.3절", "5.2절": "3.2절",
        "6.4절": "4.3절", "6.3절": "4.2절", "6.5절": "4.4절", "6.1절": "4.1절"}
RED = re.compile(r'<w:color w:val="([0-9A-Fa-f]{6})"')


def is_red(h):
    r, g, b = (int(h[i:i + 2], 16) for i in (0, 2, 4))
    return r >= 130 and g <= r * 0.55 and b <= r * 0.55


def all_paras(d):
    """본문 문단 + 표 안 문단 전부."""
    yield from d.paragraphs
    for t in d.tables:
        for row in t.rows:
            for c in row.cells:
                yield from c.paragraphs


def set_text(p, new):
    """run 이 쪼개져 있어도 안전하게 문단 전체 글을 바꾼다(첫 run 서식 유지)."""
    runs = p._p.findall(qn("w:r"))
    if not runs:
        p.add_run(new)
        return
    for r in runs[1:]:
        p._p.remove(r)
    for t in runs[0].findall(qn("w:t")):
        runs[0].remove(t)
    el = runs[0].makeelement(qn("w:t"), {})
    el.set(qn("xml:space"), "preserve")
    el.text = new
    runs[0].append(el)


def main():
    shutil.copy2(DOC, DOC.replace("CAST_압축본.docx",
                 "versions/2026-09-24/CAST_압축본_뒤처리전.docx"))
    d = docx.Document(DOC)

    # 1) 빨간 글씨 — 표 안까지
    red = 0
    for p in all_paras(d):
        for r in list(p._p.findall(qn("w:r"))):
            m = RED.search(r.xml)
            if m and is_red(m.group(1)):
                p._p.remove(r)
                red += 1
    print(f"빨간 글씨 run {red}개 추가 제거")

    # 2) 절 제목
    for p in d.paragraphs:
        t = p.text.strip()
        if t in TITLE:
            set_text(p, TITLE[t])
            print(f"  제목  {t}  →  {TITLE[t]}")

    # 3) 상호참조 — 문단 평문에서 찾아 통째로 다시 쓴다
    n = 0
    for p in all_paras(d):
        t = p.text
        if not t or "Sukprasert 등(EuroSys" in t:      # 남의 논문 절은 그대로
            continue
        new = t
        for old, rep in XREF.items():
            new = new.replace(old, rep)
        if new != t:
            set_text(p, new)
            n += 1
            print(f"  참조  {t[:46]}…")
    print(f"상호참조 {n}문단 갱신")

    d.save(DOC)

    # 검증
    import zipfile
    x = zipfile.ZipFile(DOC).read("word/document.xml").decode()
    ps = [re.sub(r"<[^>]+>", "", m.group(0)).strip()
          for m in re.finditer(r'<w:p\b[^>]*?(?<!/)>.*?</w:p>', x, re.S)]
    print("\n=== 남은 옛 번호 참조 ===")
    bad = 0
    for i, p in enumerate(ps):
        for m in re.finditer(r'(§\s?\d+(?:\.\d+)?|[567]\.\d\s?절)', p):
            print(f"  {m.group(1):<8} [{i}] {p[:64]}")
            bad += 1
    print("  (없음)" if not bad else f"  {bad}건 남음")
    print("빨간 글씨:", "잔존" if re.search(r'<w:color w:val="(EE|FF|E0|C0|D0)0000"', x) else "없음")


if __name__ == "__main__":
    main()
