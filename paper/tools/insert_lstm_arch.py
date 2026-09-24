# -*- coding: utf-8 -*-
"""[15] LSTM 구조도를 새 그림 2로 끼우고 뒤 번호를 민다.

사용자: "여기는 그래도 그 LSTM 설명할 때 저거 좋은데, 그 2층 레이어 느낌이나
그 받는 정보 13가지 정도의 그림이 나왔으면 좋겠어."

내가 블록도를 "교과서 그림"이라며 버리고 예측 곡선으로 바꾼 것이 성급했다.
곡선은 **성능**을 보이고 블록도는 **구조**를 보인다 — 둘은 대체 관계가 아니다.
구조도를 그림 2로 넣고, 예측 곡선을 그림 3으로 민다.

수치는 carbon_forecast_lstm/carbon_forecast.py 상수에서 그대로 옮겼다.
사용자가 말한 "13가지"가 실제로 맞다 — 공통 10 + 기상 3.
"""
import copy
import re
import shutil

import docx
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Pt
from PIL import Image

DOC = "/Users/jongha/Desktop/GitHub/carbon-aware-scheduler/paper/CAST_압축본.docx"
FIG = "/Users/jongha/Desktop/GitHub/carbon-aware-scheduler/paper/diagram/fig2_lstm_arch.png"
CAP = ("그림 2. 리전별 LSTM 구조 — 과거 168시간의 피처 13개(공통 10개, 기상 3개 리전만 "
       "날씨 3개 추가)를 은닉 64의 2층 LSTM에 넣고, 마지막 은닉 상태로 향후 24시간을 낸다. "
       "기상 리전은 미래 24시간 날씨 72개 값을 함께 결합한다.")


def main():
    shutil.copy2(DOC, DOC.replace("CAST_압축본.docx",
                 "versions/2026-09-24/CAST_압축본_LSTM구조도전.docx"))
    d = docx.Document(DOC)

    # ── 1) 그림 2 이상 번호를 한 칸씩 민다 (큰 번호부터) ──
    caps = {}
    for p in d.paragraphs:
        m = re.match(r"^그림\s?(\d+)\.", p.text.strip())
        if m:
            caps[int(m.group(1))] = p
    for n in sorted((k for k in caps if k >= 2), reverse=True):
        p = caps[n]
        t = p.text.strip()
        new_t = re.sub(r"^그림\s?\d+\.", f"그림 {n+1}.", t)
        runs = p._p.findall(qn("w:r"))
        for r in runs[1:]:
            p._p.remove(r)
        for e in runs[0].findall(qn("w:t")):
            runs[0].remove(e)
        el = runs[0].makeelement(qn("w:t"), {})
        el.set(qn("xml:space"), "preserve")
        el.text = new_t
        runs[0].append(el)
        print(f"  그림 {n} → 그림 {n+1}")

    # 본문 참조도 (큰 번호부터)
    for n in sorted((k for k in caps if k >= 2), reverse=True):
        for p in d.paragraphs:
            t = p.text
            if re.match(r"^그림\s?\d+\.", t.strip()) or f"그림 {n}" not in t:
                continue
            for r in p.runs:
                if f"그림 {n}" in r.text:
                    r.text = r.text.replace(f"그림 {n}", f"그림 {n+1}")
            print(f"  본문 참조 그림 {n} → {n+1}: {p.text[:42]}…")

    # ── 2) 구조도를 그림 1 캡션 뒤에 삽입 ──
    paras = d.paragraphs
    ai = next(i for i, p in enumerate(paras)
              if p.text.strip().startswith("그림 1. CAST의 전체 파이프라인"))
    anchor = paras[ai]

    im = Image.open(FIG)
    w = Pt(205)
    fp = d.add_paragraph()
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fp.add_run().add_picture(FIG, width=w)
    cp = docx.text.paragraph.Paragraph(copy.deepcopy(anchor._p), anchor._parent)
    for r in list(cp._p.findall(qn("w:r"))):
        cp._p.remove(r)
    cp.add_run(CAP).font.size = Pt(9)

    body = anchor._p.getparent()
    i = list(body).index(anchor._p)
    fp._p.getparent().remove(fp._p)
    body.insert(i + 1, fp._p)
    body.insert(i + 2, cp._p)

    d.save(DOC)
    d2 = docx.Document(DOC)
    print(f"\n그림 {sum(1 for p in d2.paragraphs if 'graphicData' in p._p.xml)}장")
    for p in d2.paragraphs:
        if re.match(r"^그림\s?\d+\.", p.text.strip()):
            print("  ", p.text.strip()[:58])


if __name__ == "__main__":
    main()
